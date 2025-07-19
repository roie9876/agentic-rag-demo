#!/usr/bin/env python3
"""
More detailed analysis of the DOCX content
"""

from docx import Document
import os

def detailed_analysis(filename):
    """Detailed analysis of DOCX content"""
    
    doc = Document(filename)
    
    print(f"📄 Document: {filename}")
    print(f"📏 File size: {os.path.getsize(filename):,} bytes")
    print(f"📝 Total paragraphs: {len(doc.paragraphs)}")
    print(f"📑 Sections: {len(doc.sections)}")
    
    # Look for any indicators of page structure
    page_breaks = 0
    manual_breaks = 0
    empty_paras = 0
    
    for i, para in enumerate(doc.paragraphs):
        text = para.text
        if '\x0c' in text:  # Form feed
            page_breaks += 1
        if '\n\n' in text or text.strip() == '':
            if text.strip() == '':
                empty_paras += 1
        if 'עמוד' in text or 'page' in text.lower():  # Hebrew word for "page"
            manual_breaks += 1
            
    print(f"🔀 Page breaks found: {page_breaks}")
    print(f"📄 Manual page references: {manual_breaks}")
    print(f"⬜ Empty paragraphs: {empty_paras}")
    
    # Show content distribution
    print(f"\n📊 Content sample (every 50th paragraph):")
    for i in range(0, min(len(doc.paragraphs), 350), 50):
        text = doc.paragraphs[i].text[:80] + '...' if len(doc.paragraphs[i].text) > 80 else doc.paragraphs[i].text
        if text.strip():
            print(f"   Para {i+1}: '{text}'")

if __name__ == "__main__":
    detailed_analysis("iecdocsmall.docx")
